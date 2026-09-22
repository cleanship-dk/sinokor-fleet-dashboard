"""Regression tests for the ways a *new* Excel dump can differ from the one we built against.

Each test mutates a copy of the shipped workbook and asserts that the dashboard either still
gets the right answer, or says loudly that it cannot. The failure mode these guard against is
the dangerous one: a dump that parses "successfully" and quietly produces a wrong list.
"""
from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from cls.brief import build_email  # noqa: E402
from cls.loader import load_workbook  # noqa: E402
from cls.names import build_index, match  # noqa: E402
from cls.rules import PRIME, Settings, evaluate, month_window  # noqa: E402
from cls.store import load_watchlist, save_upload  # noqa: E402

TODAY = _dt.date(2026, 9, 22)
# The client workbook is deliberately not in the repo (see .gitignore), so on a fresh
# clone these tests skip with an explanation rather than erroring out.
_WORKBOOKS = sorted((ROOT / "data").glob("*.xlsx"))
WORKBOOK = _WORKBOOKS[0] if _WORKBOOKS else None

pytestmark = pytest.mark.skipif(
    WORKBOOK is None,
    reason="No fleet workbook in data/ - drop an .xlsx export there to run these tests.",
)

EXPECTED_PRIME = {
    "ATLANTIC SOUTH", "SAWASDEE VEGA", "SAWASDEE MIMOSA", "INCHEON VOYAGER",
    "SINOKOR NIIGATA", "YEOSU VOYAGER", "HEUNG-A YOUNG", "TOYAMA TRADER",
}


@pytest.fixture
def settings():
    start, end = month_window(2026, 10)
    return Settings(today=TODAY, window_start=start, window_end=end, window_label="October 2026")


@pytest.fixture
def mutate(tmp_path):
    """Open a copy of the workbook, let the test edit it, and return the reloaded result."""
    def _mutate(edit, name="mutated.xlsx"):
        book = openpyxl.load_workbook(WORKBOOK)
        edit(book)
        path = tmp_path / name
        book.save(path)
        return load_workbook(path)
    return _mutate


def primes(wb, settings):
    result = evaluate(wb, settings)
    return set(result.loc[result["status"] == PRIME, "vessel"])


# ---------------------------------------------------------------------------- sheet naming


def test_a_renamed_fleet_sheet_is_still_found_by_its_contents(mutate, settings):
    """'fleet' as a loose alias used to match 'Fleet Contact Info', which parses to 73 rows of
    the wrong thing and silently yields zero candidates."""
    wb = mutate(lambda b: setattr(b["Fleet Schedule"], "title", "Schedule 2026"))
    assert wb.sheet_map["fleet"] == "Schedule 2026"
    assert wb.sheet_map["contacts"] == "Fleet Contact Info"
    assert primes(wb, settings) == EXPECTED_PRIME


@pytest.mark.parametrize("new_name", ["Vessels", "Master list", "2026 Fleet", "FLEET SCHEDULE(10)"])
def test_the_fleet_sheet_is_found_whatever_it_is_called(mutate, settings, new_name):
    wb = mutate(lambda b: setattr(b["Fleet Schedule"], "title", new_name))
    assert wb.sheet_map["fleet"] == new_name
    assert primes(wb, settings) == EXPECTED_PRIME


def test_no_sheet_ever_serves_two_roles(mutate, settings):
    wb = mutate(lambda b: None)
    named = [v for v in wb.sheet_map.values() if v]
    assert len(named) == len(set(named))


def test_a_workbook_with_no_fleet_sheet_fails_loudly(tmp_path):
    book = openpyxl.Workbook()
    book.active.title = "Random"
    book.active["A1"] = "nothing useful"
    path = tmp_path / "useless.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="fleet schedule"):
        load_workbook(path)


# ---------------------------------------------------------------------------- headers


@pytest.mark.parametrize("spelling", [
    "Last DryDocking", "Last Dry-Docking", "LAST DRY DOCKING", "Last  Dry  Docking", "Last DD",
])
def test_drydock_header_spellings_all_map(mutate, settings, spelling):
    """An unmapped drydock column used to push all 73 vessels onto the contact-sheet fallback,
    which made NINGBO TRADER - the vessel that must never qualify - a prime candidate."""
    wb = mutate(lambda b: b["Fleet Schedule"].cell(3, 8).__setattr__("value", spelling))
    result = evaluate(wb, settings)
    assert (result["drydock_source"] == "Fleet Schedule").sum() == 64
    assert result.loc[result["vessel"] == "NINGBO TRADER", "status"].iloc[0] != PRIME
    assert set(result.loc[result["status"] == PRIME, "vessel"]) == EXPECTED_PRIME


def test_port_call_columns_may_be_reordered(mutate, settings):
    """Hub used to have to come first: the loader only attached eta/etd to a block a hub had
    already opened, so eta/etd/hub order silently dropped calls."""
    def rotate(book):
        sheet = book["Fleet Schedule"]
        for row in range(3, 77):
            hub, term, eta, etd = (sheet.cell(row, c).value for c in (19, 20, 21, 22))
            for col, value in zip((19, 20, 21, 22), (term, eta, etd, hub)):
                sheet.cell(row, col).value = value

    wb = mutate(rotate)
    assert len(wb.calls) == 96
    assert primes(wb, settings) == EXPECTED_PRIME


def test_a_third_port_call_block_is_read(mutate, settings):
    def add_block(book):
        sheet = book["Fleet Schedule"]
        for col, header in zip((28, 29, 30, 31), ("Next CLS Hub", "Terminal", "ETA", "ETD")):
            sheet.cell(3, col).value = header
        sheet.cell(8, 28).value = "Singapore"           # HEUNG-A YOUNG gains a third call
        sheet.cell(8, 29).value = "PSA"
        sheet.cell(8, 30).value = _dt.datetime(2026, 10, 27)
        sheet.cell(8, 31).value = _dt.datetime(2026, 10, 28)

    wb = mutate(add_block)
    result = evaluate(wb, settings)
    assert result.loc[result["vessel"] == "HEUNG-A YOUNG", "calls_in_window"].iloc[0] == 3


# ---------------------------------------------------------------------------- dates


def test_an_undecidable_date_column_is_flagged(mutate):
    """If nothing in a text date column settles day-first vs month-first, every date in it can
    be wrong by up to eleven months. Guessing silently is the unacceptable outcome."""
    def make_ambiguous(book):
        sheet = book["Fleet Schedule"]
        for row in range(4, 77):
            for col in (21, 22, 25, 26):
                value = sheet.cell(row, col).value
                if hasattr(value, "strftime"):
                    if value.day > 12:
                        sheet.cell(row, col).value = None      # would have settled it
                    else:
                        sheet.cell(row, col).value = value.strftime("%m/%d/%Y")
                elif isinstance(value, str):
                    sheet.cell(row, col).value = None

    wb = mutate(make_ambiguous)
    flagged = [i for i in wb.issues if "settles" in i.detail]
    assert {i.field for i in flagged} == {"call1.eta", "call1.etd", "call2.eta", "call2.etd"}


def test_a_column_that_proves_its_own_convention_is_not_flagged(mutate):
    wb = mutate(lambda b: None)
    assert not [i for i in wb.issues if "settles" in i.detail]


# ---------------------------------------------------------------------------- missing sheets


@pytest.mark.parametrize("sheet,field", [("Order Log", "orders"), ("Invoice", "invoices")])
def test_a_missing_history_sheet_is_reported(mutate, sheet, field):
    """Losing either sheet makes vessels look more overdue than they are, so it cannot be silent."""
    wb = mutate(lambda b: b.remove(b[sheet]))
    assert [i for i in wb.issues if i.field == field]


def test_a_non_numeric_loa_is_reported(mutate):
    def spoil(book):
        sheet = book["Fleet Schedule"]
        for row in range(4, 77):
            value = sheet.cell(row, 7).value
            if value is not None:
                sheet.cell(row, 7).value = "%s m" % value

    wb = mutate(spoil)
    assert len([i for i in wb.issues if i.field == "loa"]) > 60


def test_a_duplicated_vessel_row_is_collapsed_and_reported(mutate, settings):
    def duplicate(book):
        sheet = book["Fleet Schedule"]
        for col in range(1, 28):
            sheet.cell(77, col).value = sheet.cell(8, col).value   # a second HEUNG-A YOUNG

    wb = mutate(duplicate)
    assert (wb.fleet["vessel"] == "HEUNG-A YOUNG").sum() == 1
    assert [i for i in wb.issues if i.field == "vessel" and i.vessel == "HEUNG-A YOUNG"]
    assert primes(wb, settings) == EXPECTED_PRIME


# ---------------------------------------------------------------------------- joins


def test_a_missing_vessel_does_not_have_its_history_given_to_a_similar_one():
    """HAKATA VOYAGER and JAKARTA VOYAGER score 0.889 on difflib. If JAKARTA drops out of a
    dump, its cleaning must not be credited to HAKATA."""
    index = build_index(["HAKATA VOYAGER", "HAKATA EXPRESS", "NINGBO TRADER"])
    assert match("JAKARTA VOYAGER", index) is None
    assert match("Jakarta Voyager", index) is None


@pytest.mark.parametrize("written,canonical", [
    ("Toyoma Trader", "TOYAMA TRADER"),
    ("Volostchny Voyager", "VOSTOCHNY VOYAGER"),
    ("Hakata Expresss", "HAKATA EXPRESS"),
    ("Pacific Nongbo", "PACIFIC NINGBO"),
])
def test_real_misspellings_still_join(written, canonical):
    index = build_index(["TOYAMA TRADER", "VOSTOCHNY VOYAGER", "HAKATA EXPRESS",
                         "PACIFIC NINGBO", "HAKATA VOYAGER", "JAKARTA VOYAGER"])
    assert match(written, index) == canonical


# ---------------------------------------------------------------------------- the client brief


def test_the_brief_repeats_the_alerts_the_dashboard_raised(settings):
    """INCHEON VOYAGER is a prime candidate that drydocks 25 Oct. Pitching it a cleaning for
    10 Oct without saying so is the kind of thing the client's fleet department spots first."""
    wb = load_workbook(WORKBOOK)
    result = evaluate(wb, settings)
    row = result[result["vessel"] == "INCHEON VOYAGER"].iloc[0].to_dict()
    email = build_email([row], settings)
    assert "Please note      : Next drydock in" in email


def test_the_brief_does_not_call_a_non_candidate_due(settings):
    wb = load_workbook(WORKBOOK)
    result = evaluate(wb, settings)
    row = result[result["vessel"] == "NINGBO TRADER"].iloc[0].to_dict()
    email = build_email([row], settings)
    assert "are due for a hull cleaning and call" not in email
    assert "raised for discussion only" in email


def test_the_brief_describes_the_thresholds_actually_in_force(settings):
    wb = load_workbook(WORKBOOK)
    row = evaluate(wb, settings).iloc[0].to_dict()
    loose = Settings(**{**settings.__dict__, "min_dd_age_years": 0.25})
    assert "0 years after drydocking" not in build_email([row], loose)


def test_a_call_that_crosses_a_month_shows_both_months(settings):
    wb = load_workbook(WORKBOOK)
    row = evaluate(wb, settings).iloc[0].to_dict()
    row["window_calls"] = [{"hub": "Busan", "terminal": "BCT",
                            "eta": _dt.date(2026, 10, 30), "etd": _dt.date(2026, 11, 3)}]
    assert "30 Oct-03 Nov" in build_email([row], settings)


# ---------------------------------------------------------------------------- persistence


@pytest.mark.parametrize("content", ['"hello"', "42", "null", "[1, 2]", "{oops"])
def test_a_corrupt_watchlist_never_takes_the_app_down(tmp_path, monkeypatch, content):
    import cls.store as store
    path = tmp_path / "watchlist.json"
    path.write_text(content, encoding="utf-8")
    monkeypatch.setattr(store, "WATCHLIST_FILE", path)
    assert isinstance(store.load_watchlist(), dict)


def test_two_uploads_in_the_same_second_do_not_overwrite_each_other(tmp_path, monkeypatch):
    import cls.store as store
    monkeypatch.setattr(store, "DATA_DIR", tmp_path)
    first, _ = store.save_upload(b"AAAA", "dump.xlsx")
    second, _ = store.save_upload(b"BBBB", "dump.xlsx")
    assert first != second
    assert first.read_bytes() == b"AAAA"
    assert second.read_bytes() == b"BBBB"
